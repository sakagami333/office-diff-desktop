import sys
from pathlib import Path
from docx import Document
from pptx import Presentation
from pptx.util import Inches, Pt

out = Path(sys.argv[1])
version = sys.argv[2]  # v1 or v2
out.mkdir(parents=True, exist_ok=True)
v2 = version == 'v2'

# ---- Word: 研修運営マニュアル ----
d = Document()
d.add_heading('研修運営マニュアル', 0)
d.add_heading('1. 目的', 1)
d.add_paragraph('本マニュアルは、ラーニングサポート部が実施する研修の運営手順を定めるものです。')
d.add_heading('2. 事前準備', 1)
d.add_paragraph('研修開始の3営業日前までに、以下を準備します。' if not v2 else
                '研修開始の5営業日前までに、以下を準備します。')
for item in ['受講者名簿の確認', '教材の印刷', '演習環境の動作確認'] + (['受講者への事前案内メール送信'] if v2 else []):
    d.add_paragraph(item, style='List Bullet')
d.add_heading('3. 当日の流れ', 1)
t = d.add_table(rows=1, cols=2)
t.style = 'Table Grid'
t.rows[0].cells[0].text, t.rows[0].cells[1].text = '時間', '内容'
rows = [('9:30', '受付開始'), ('10:00', '講義開始'), ('12:00', '昼休憩'), ('17:00', '終了・アンケート')]
if v2:
    rows = [('9:15', '受付開始'), ('9:45', 'オリエンテーション'), ('10:00', '講義開始'),
            ('12:00', '昼休憩'), ('17:00', '終了・アンケート')]
for a, b in rows:
    c = t.add_row().cells
    c[0].text, c[1].text = a, b
d.add_heading('4. 事後対応', 1)
d.add_paragraph('アンケート結果は翌営業日までに集計し、部内で共有します。')
if not v2:
    d.add_paragraph('欠席者には後日、資料をメールで送付します。')
else:
    d.add_paragraph('欠席者には後日、資料と録画URLをメールで送付します。')
d.save(out / '研修運営マニュアル.docx')

# ---- PowerPoint: 研修案内 ----
p = Presentation()
s = p.slides.add_slide(p.slide_layouts[0])
s.shapes.title.text = 'Python基礎研修のご案内'
s.placeholders[1].text = 'ラーニングサポート部　2026年11月' if not v2 else 'ラーニングサポート部　2026年12月'

s = p.slides.add_slide(p.slide_layouts[1])
s.shapes.title.text = '研修概要'
tf = s.placeholders[1].text_frame
items = [('対象：入社1〜2年目のエンジニア', 0), ('期間：2日間', 0), ('形式：対面（東京会場）', 0)]
if v2:
    items = [('対象：入社1〜3年目のエンジニア', 0), ('期間：2日間', 0),
             ('形式：対面（東京会場）／オンライン併用', 0), ('定員：20名', 0)]
tf.text = items[0][0]
for text, lvl in items[1:]:
    para = tf.add_paragraph(); para.text = text; para.level = lvl
s.notes_slide.notes_text_frame.text = '対象者の範囲は人事部と調整済み。' if not v2 else '対象者の範囲は人事部と再調整済み（3年目まで拡大）。'

s = p.slides.add_slide(p.slide_layouts[1])
s.shapes.title.text = 'カリキュラム'
tf = s.placeholders[1].text_frame
tf.text = '1日目：Python文法の基礎'
for text, lvl in [('変数・制御構文', 1), ('関数', 1), ('2日目：実践演習', 0), ('ファイル操作', 1)]:
    para = tf.add_paragraph(); para.text = text; para.level = lvl
if v2:
    para = tf.add_paragraph(); para.text = 'Webスクレイピング入門'; para.level = 1

s = p.slides.add_slide(p.slide_layouts[1])
s.shapes.title.text = '持ち物'
s.placeholders[1].text_frame.text = 'ノートPC（貸出あり）'

s = p.slides.add_slide(p.slide_layouts[1])
s.shapes.title.text = 'お申し込み'
s.placeholders[1].text_frame.text = '社内ポータルの研修申込フォームから、10月31日までにお申し込みください。' if not v2 else \
    '社内ポータルの研修申込フォームから、11月14日までにお申し込みください。'
if v2:
    # 「持ち物」スライドを削除。PowerPointでの削除と同じく、他のスライドのIDは変わらない
    sld_ids = p.slides._sldIdLst
    target = list(sld_ids)[3]
    p.part.drop_rel(target.rId)
    sld_ids.remove(target)
p.save(out / '研修案内.pptx')

# ---- Excel: 研修実績.xlsx ----
# v1 → v2 の変更: 値の変更 / 行の挿入 / 数式と計算結果の変化 / 日付 / シート名の変更 / シートの追加・削除
import datetime
import re
import zipfile
from openpyxl import Workbook

wb = Workbook()
ws = wb.active
ws.title = '実績'
ws.append(['研修名', '開催日', '受講者数', '満足度'])
rows = [('Python基礎', datetime.date(2026, 9, 10), 18, 4.5),
        ('Linux入門', datetime.date(2026, 9, 24), 12, 4.2),
        ('クラウド概論', datetime.date(2026, 10, 8), 20, 4.7)]
if v2:
    rows = [('Python基礎', datetime.date(2026, 9, 10), 18, 4.5),
            ('Linux入門', datetime.date(2026, 9, 24), 15, 4.2),       # 受講者数 12 → 15
            ('Git/GitHub実践', datetime.date(2026, 10, 1), 16, 4.6),  # 行を挿入
            ('クラウド概論', datetime.date(2026, 10, 8), 20, 4.7)]
for r in rows:
    ws.append(list(r))
    ws.cell(row=ws.max_row, column=2).number_format = 'yyyy/mm/dd'
total_row = ws.max_row + 1
ws.cell(row=total_row, column=1, value='合計')
ws.cell(row=total_row, column=3, value=f'=SUM(C2:C{total_row - 1})')
total = sum(r[2] for r in rows)

ws2 = wb.create_sheet('講師一覧' if not v2 else '講師')   # シート名の変更（中身は同じ）
for r in [('氏名', '担当'), ('山田', 'Python'), ('佐藤', 'Linux'), ('鈴木', 'クラウド')]:
    ws2.append(list(r))
if not v2:
    ws3 = wb.create_sheet('旧メモ')                          # v2 では削除
    ws3.append(['このシートは次の版で削除'])
else:
    ws3 = wb.create_sheet('来期予定')                        # v2 で追加
    ws3.append(['研修名', '予定月'])
    ws3.append(['AWS入門', '2027年1月'])
xlsx = out / '研修実績.xlsx'
wb.save(xlsx)

# openpyxl は数式の計算結果を保存しないため、Excel で保存したファイルと同じく結果（<v>）を書き込む
with zipfile.ZipFile(xlsx) as z:
    items = {n: z.read(n) for n in z.namelist()}
sheet = items['xl/worksheets/sheet1.xml'].decode('utf-8')
sheet, n = re.subn(r'(<f>SUM\([^<]*\)</f>)(<v\s*/>|<v></v>)?', rf'\1<v>{total}</v>', sheet)
assert n == 1, n
items['xl/worksheets/sheet1.xml'] = sheet.encode('utf-8')
with zipfile.ZipFile(xlsx, 'w', zipfile.ZIP_DEFLATED) as z:
    for name, data in items.items():
        z.writestr(name, data)
print('done', version)
